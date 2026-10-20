"""Build the rater, judge, and compiled spreadsheets.

Inputs : 0_Data/citation_map.json
         2_Output/agent_work/raters_merged.json, disagreements.json, agreement_stats.json  (from 04)
         2_Output/agent_work/judge_batch*.json                                           (judge agent)
Outputs: 2_Output/Rater1_coding.xlsx, Rater2_coding.xlsx, Rater3_coding.xlsx
         2_Output/Judge_adjudication.xlsx
         2_Output/Compiled_evidence_coding.xlsx

Requires openpyxl. Run: python3 -I 1_Scripts/05_build_spreadsheets.py
"""
import json
import random
import re
import statistics
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / "2_Output/agent_work"
OUT = ROOT / "2_Output"
SEED = 20261013
REC_ORDER = ["R1", "R2", "R3", "R4", "R5", "GPP1", "R6", "R7", "R8", "R9", "R10", "R11", "R12", "R13", "R14", "GPP2", "GPP3"]
COMPARATIVE = ["RCT", "Sibling/within-family", "Observational – adjusted, incl. severity",
               "Observational – adjusted, excl. severity", "Quasi-experimental (other)"]
TYPE_COLS = COMPARATIVE[:3] + ["Observational – adjusted, excl. severity", "Observational – adjusted, severity unclear",
                               "Quasi-experimental (other)", "Observational – adjustment unclear", "Observational – unadjusted/descriptive", "Systematic review/meta-analysis",
                           "Pharmacokinetic/physiologic", "Case report/case series", "Survey/qualitative/program evaluation",
                           "Narrative review/commentary", "Guideline/consensus statement", "Protocol/design paper",
                           "Not a study (resource)", "Other"]
RATER_FIELDS = ["evidence_basis", "study_type", "design_detail", "severity_adjustment", "severity_measure",
                "sibling_analysis_any", "comparison_group", "population", "perinatal_population",
                "review_component_designs", "n_studies", "sample_size_n", "sample_size_unit", "sample_size_detail",
                "sample_size_source", "confidence", "rationale"]
FINAL_FIELDS = ["study_type", "severity_adjustment", "severity_measure", "sibling_analysis_any",
                "sample_size_n", "sample_size_unit", "n_studies", "sample_size_detail"]
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
FLAG_FILL = PatternFill("solid", fgColor="FFF2CC")
INPUT_FILL = PatternFill("solid", fgColor="E2EFDA")


def sort_key(rid):
    return (0, int(rid)) if rid.isdigit() else (1, rid)


def write_sheet(ws, header, rows, widths=None, flag_col=None, input_cols=()):
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="top")
    for r in rows:
        ws.append([", ".join(map(str, v)) if isinstance(v, list) else
                   "; ".join(f"{k}: {x}" for k, x in v.items()) if isinstance(v, dict) else v for v in r])
    widths = widths or {}
    for i, h in enumerate(header, 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14 if len(h) < 14 else min(len(h) + 2, 40))
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if flag_col is not None and row[flag_col].value == "Yes":
            for c in row:
                c.fill = FLAG_FILL
        for j in input_cols:
            row[j].fill = INPUT_FILL
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


WIDE = {"citation": 60, "title": 45, "rationale": 60, "design_detail": 40, "severity_measure": 40,
        "comparison_group": 30, "population": 30, "sample_size_detail": 35, "judge_rationale": 70,
        "flag_reasons": 50, "rec_text": 70, "human_notes": 40, "paper_file": 40, "decision_by_field": 60}


def title_match(c):
    """Share of PubMed-title words (>3 letters) found in the ACOG citation: a check on PMID resolution."""
    if not c.get("title"):
        return None
    norm = lambda s: re.sub(r"[^a-z0-9 ]", " ", s.lower().replace("-", ""))
    words = [w for w in norm(c["title"]).split() if len(w) > 3]
    cit = norm(c["citation"])
    return round(sum(w in cit for w in words) / len(words), 2) if words else None


def main():
    cmap = json.loads((ROOT / "0_Data/citation_map.json").read_text())
    recs, cits = cmap["recs"], {c["ref_id"]: c for c in cmap["citations"]}
    merged = json.loads((WORK / "raters_merged.json").read_text())
    dis = {d["ref_id"]: d for d in json.loads((WORK / "disagreements.json").read_text())}
    stats = json.loads((WORK / "agreement_stats.json").read_text())
    judge = {}
    for f in sorted(WORK.glob("judge_batch*.json")) + sorted(WORK.glob("judge_round2.json")):  # round 2 wins
        for row in json.loads(f.read_text()):
            judge[str(row["ref_id"])] = row
    spot = {}
    if (WORK / "judge_spotcheck_results.json").exists():
        spot = {str(r["ref_id"]): r for r in json.loads((WORK / "judge_spotcheck_results.json").read_text())}
    # spot-checks judged round-1 codes; drop them for citations re-coded in round 2
    spot = {k: {**v, "round": 1} for k, v in spot.items() if merged[k]["rater1"].get("coding_round", 1) == 1}
    if (WORK / "judge_spotcheck_round2_results.json").exists():
        for r in json.loads((WORK / "judge_spotcheck_round2_results.json").read_text()):
            spot[str(r["ref_id"])] = {**r, "round": 2}
    overrides = {}
    if (WORK / "pi_overrides.json").exists():
        overrides = json.loads((WORK / "pi_overrides.json").read_text())
    rids = sorted(cits, key=sort_key)
    base_cols = ["ref_id", "recs", "citation", "pmid", "pmcid", "doi", "full_text_status", "paper_file"]
    base = lambda c: [c["ref_id"], c["recs"], c["citation"], c.get("pmid"), c.get("pmcid"), c.get("doi"),
                      c.get("full_text_status"), c.get("paper_file")]

    # ---------- per-rater workbooks ----------
    for r in (1, 2, 3):
        wb = Workbook()
        ws = wb.active
        ws.title = f"Rater{r}"
        rows = [base(cits[rid]) + [merged[rid].get(f"rater{r}", {}).get(f) for f in RATER_FIELDS] for rid in rids]
        write_sheet(ws, base_cols + RATER_FIELDS, rows, WIDE)
        wb.save(OUT / f"Rater{r}_coding.xlsx")

    # ---------- judge workbook ----------
    wb = Workbook()
    ws = wb.active
    ws.title = "Adjudications"
    jcols = ["ref_id", "recs", "citation", "evidence_basis", "disputed_fields"] + \
            [f"rater{r}_{f}" for f in ("study_type", "severity_adjustment", "sibling_analysis_any", "sample_size_n") for r in (1, 2, 3)] + \
            [f"final_{f}" for f in FINAL_FIELDS] + ["decision_by_field", "judge_rationale", "judge_confidence", "flag_for_human", "flag_reason"]
    jrows = []
    for rid in sorted(dis, key=sort_key):
        d, j, c = dis[rid], judge.get(rid, {}), cits[rid]
        rv = d["rater_codes"]
        jrows.append([rid, c["recs"], c["citation"], d.get("evidence_basis"), d["disputed_fields"]] +
                     [(rv.get(f"rater{r}") or {}).get(f) for f in ("study_type", "severity_adjustment", "sibling_analysis_any", "sample_size_n") for r in (1, 2, 3)] +
                     [j.get(f"final_{f}") for f in FINAL_FIELDS] +
                     [j.get("decision_by_field"), j.get("judge_rationale"), j.get("judge_confidence"), j.get("flag_for_human"), j.get("flag_reason")])
    write_sheet(ws, jcols, jrows, WIDE, flag_col=jcols.index("flag_for_human"))
    ws = wb.create_sheet("Spot_check_unanimous")
    scols = ["ref_id", "round", "recs", "citation", "study_type", "severity_adjustment", "sibling_analysis_any", "sample_size_n",
             "spotcheck_agrees", "concern", "flag_for_human", "flag_reason"]
    srows = []
    for rid in sorted(spot, key=sort_key):
        r1 = merged[rid]["rater1"]
        srows.append([rid, spot[rid]["round"], cits[rid]["recs"], cits[rid]["citation"], r1.get("study_type"), r1.get("severity_adjustment"),
                      r1.get("sibling_analysis_any"), r1.get("sample_size_n")] +
                     [spot[rid].get(k) for k in ("spotcheck_agrees", "concern", "flag_for_human", "flag_reason")])
    write_sheet(ws, scols, srows, {**WIDE, "concern": 60, "flag_reason": 50}, flag_col=scols.index("flag_for_human"))
    wb.save(OUT / "Judge_adjudication.xlsx")

    # ---------- compiled: final value per citation ----------
    final = {}
    for rid in rids:
        rv = merged[rid]
        codes = [rv.get(f"rater{r}", {}) for r in (1, 2, 3)]
        d = dis.get(rid)
        j = judge.get(rid) if d else None  # judge rulings apply only to current disagreements
        disputed = d["disputed_fields"] if d else []
        f = {}
        maj = Counter(x.get("study_type") for x in codes).most_common(1)[0]
        f["study_type"] = j.get("final_study_type") if j and "study_type" in disputed else maj[0]
        ref_rater = next((x for x in codes if x.get("study_type") == f["study_type"]), codes[0])
        for k in ("severity_adjustment", "sibling_analysis_any"):
            f[k] = j.get(f"final_{k}") if j and k in disputed else codes[0].get(k)
        sizes = [x.get("sample_size_n") for x in codes if x.get("sample_size_n") is not None]
        f["sample_size_n"] = j.get("final_sample_size_n") if j and "sample_size_n" in disputed else (int(statistics.median(sizes)) if sizes else None)
        ns = [x.get("n_studies") for x in codes if x.get("n_studies") is not None]
        f["n_studies"] = j.get("final_n_studies") if j and "n_studies" in disputed else (int(statistics.median(ns)) if ns else None)
        for k in ("severity_measure", "sample_size_unit", "sample_size_detail"):
            f[k] = (j.get(f"final_{k}") if j and j.get(f"final_{k}") else None) or ref_rater.get(k)
        for k in ("design_detail", "comparison_group", "population", "perinatal_population", "review_component_designs"):
            f[k] = ref_rater.get(k)
        f["evidence_basis"] = "citation_only" if not cits[rid].get("pmid") else ("full_text" if cits[rid]["full_text_available"] else "abstract_only")
        f["has_abstract"] = "No" if "_No abstract in PubMed._" in (ROOT / cits[rid]["paper_file"]).read_text() or not cits[rid].get("pmid") else "Yes"
        st = f["study_type"]
        if st == "Observational – adjusted, excl. severity" and f["severity_adjustment"] == "Unclear":
            st = "Observational – adjusted, severity unclear"
        elif st == "Observational – unadjusted/descriptive" and f["severity_adjustment"] == "Unclear":
            st = "Observational – adjustment unclear"
        f["study_type_detailed"] = st
        # flags for human review
        reasons = []
        if j and str(j.get("flag_for_human", "")).lower().startswith("y"):
            reasons.append("Judge: " + (j.get("flag_reason") or "flagged"))
        sp = spot.get(rid)
        if sp and str(sp.get("spotcheck_agrees", "")).lower().startswith("n"):
            reasons.append("Judge spot-check disputes unanimous raters: " + (sp.get("concern") or ""))
        elif sp and str(sp.get("flag_for_human", "")).lower().startswith("y"):
            reasons.append("Judge spot-check: " + (sp.get("flag_reason") or "flagged"))
        f["judge_spotcheck"] = ("" if not sp else ("Agrees" if str(sp.get("spotcheck_agrees", "")).lower().startswith("y") else "Disputes"))
        if d and d.get("three_way_split_study_type"):
            reasons.append("3-way rater split on study type")
        if j and "study_type" in disputed and maj[1] >= 2 and j.get("final_study_type") != maj[0]:
            reasons.append("Judge overrode 2-rater majority on study type")
        if not d and any(x.get("confidence") == "Low" for x in codes):
            reasons.append("Unanimous, but a rater reported Low confidence")
        if f["evidence_basis"] != "full_text" and (f["study_type"] in COMPARATIVE[1:5] or
                                                   (f["study_type"] == "Observational – unadjusted/descriptive" and f["severity_adjustment"] == "Unclear")):
            reasons.append("No full text: adjustment/severity coded from abstract, verify against full text")
        tm = title_match(cits[rid])
        if tm is not None and tm < 0.6:
            reasons.append(f"PubMed title only partly matches ACOG citation (overlap {tm})")
        ov = overrides.get(rid, {})
        for k, v in ov.items():
            if not k.startswith("_"):
                f[k] = v
        f["override"] = (f"{', '.join(k for k in ov if not k.startswith('_'))} — {ov.get('_reason', '')}") if ov else ""
        f["agreement"] = "Unanimous" if not d else ("Adjudicated: " + ", ".join(disputed))
        f["coding_round"] = codes[0].get("coding_round", 1)
        f["flag"] = "Yes" if reasons else "No"
        f["flag_reasons"] = "; ".join(reasons)
        f["title_match"] = tm
        f["rater_types"] = " | ".join(f"R{r}: {x.get('study_type')}" for r, x in zip((1, 2, 3), codes))
        f["rater_sizes"] = " | ".join(f"R{r}: {x.get('sample_size_n')}" for r, x in zip((1, 2, 3), codes))
        final[rid] = f

    wb = Workbook()
    # README
    ws = wb.active
    ws.title = "README"
    readme = [
        ["ACOG CPG No. 5 — evidence coding (compiled)"],
        ["Each citation used in a recommendation's evidence summary was coded independently by 3 AI raters against coding_protocol.md; "
         "disagreements on study type, severity adjustment, sibling analysis, sample size (>10% apart), or number of review studies were adjudicated by an AI judge. "
         "The judge also spot-checked (flag only, no changes) every unanimously coded comparative study."],
        ["Caution: the 3 raters are instances of the same model reading the same files, so their errors are correlated; agreement statistics below overstate reliability relative to independent human coders."],
        ["Final value = unanimous rater value, or the judge's value for disputed fields. Sample size for agreeing raters = median of the three."],
        ["flag = Yes rows (yellow) need human review; see flag_reasons. Green columns in Human_verification are for the human reviewer."],
        ["See METHODS_AND_VERIFICATION.md in the project root for methods and the verification plan."],
        [],
        ["Agreement statistic", "Value"],
        ["Citations coded", stats["n_citations"]],
        ["Citations with any disagreement", stats["n_with_any_disagreement"]],
    ] + [[f"% agreement — {k}", v] for k, v in stats["percent_agreement"].items()] + \
        [[f"Fleiss' kappa — {k}", v] for k, v in stats["fleiss_kappa"].items()]
    for row in readme:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=14)
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 14
    for r in range(2, 7):
        ws.cell(r, 1).alignment = Alignment(wrap_text=True)

    # Citations_final
    ws = wb.create_sheet("Citations_final")
    fcols = ["ref_id", "recs", "citation", "title", "pmid", "pmcid", "doi", "full_text_status", "evidence_basis",
             "has_abstract", "study_type", "study_type_detailed", "design_detail", "severity_adjustment", "severity_measure", "sibling_analysis_any",
             "comparison_group", "population", "perinatal_population", "review_component_designs", "n_studies",
             "sample_size_n", "sample_size_unit", "sample_size_detail", "agreement", "coding_round", "judge_spotcheck", "override", "rater_types", "rater_sizes",
             "flag", "flag_reasons", "title_match", "acog_citation_error", "paper_file"]
    rows = []
    for rid in rids:
        c, f = cits[rid], final[rid]
        src = {**c, **f}
        rows.append([src.get(k) for k in fcols])
    write_sheet(ws, fcols, rows, WIDE, flag_col=fcols.index("flag"))

    # Rec_x_Citation (long)
    ws = wb.create_sheet("Rec_x_Citation")
    lcols = ["rec_id", "section", "strength", "evidence_quality", "rec_text", "ref_id", "citation", "study_type_detailed",
             "severity_adjustment", "sibling_analysis_any", "sample_size_n", "sample_size_unit", "n_studies",
             "evidence_basis", "flag"]
    rows = []
    for k in REC_ORDER:
        rr = recs[k]
        for rid in rr["citations"]:
            f = final[rid]
            rows.append([k, rr["section"], rr["strength"], rr["evidence_quality"], rr["text"], rid, cits[rid]["citation"],
                         f["study_type_detailed"], f["severity_adjustment"], f["sibling_analysis_any"], f["sample_size_n"],
                         f["sample_size_unit"], f["n_studies"], f["evidence_basis"], f["flag"]])
    write_sheet(ws, lcols, rows, WIDE, flag_col=lcols.index("flag"))

    # Rec_summary
    ws = wb.create_sheet("Rec_summary")
    scols = ["rec_id", "section", "strength", "evidence_quality", "rec_text", "n_citations", "n_full_text"] + TYPE_COLS + \
            ["n_any_sibling_analysis", "median_N_comparative_studies", "max_N_comparative_studies", "n_flagged"]
    rows = []
    for k in REC_ORDER:
        rr = recs[k]
        fs = [final[rid] for rid in rr["citations"]]
        tc = Counter(f["study_type_detailed"] for f in fs)
        comp_n = [f["sample_size_n"] for f in fs if f["study_type"] in COMPARATIVE and f["sample_size_n"]]
        rows.append([k, rr["section"], rr["strength"], rr["evidence_quality"], rr["text"], len(fs),
                     sum(f["evidence_basis"] == "full_text" for f in fs)] + [tc.get(t, 0) for t in TYPE_COLS] +
                    [sum(f["sibling_analysis_any"] == "Yes" for f in fs),
                     int(statistics.median(comp_n)) if comp_n else None, max(comp_n) if comp_n else None,
                     sum(f["flag"] == "Yes" for f in fs)])
    write_sheet(ws, scols, rows, {**WIDE, **{t: 13 for t in TYPE_COLS}})

    # Rater_comparison
    ws = wb.create_sheet("Rater_comparison")
    ccols = ["ref_id"] + [f"rater{r}_{f}" for f in ("study_type", "severity_adjustment", "sibling_analysis_any", "sample_size_n", "confidence") for r in (1, 2, 3)] + ["agreement"]
    rows = [[rid] + [merged[rid].get(f"rater{r}", {}).get(f) for f in ("study_type", "severity_adjustment", "sibling_analysis_any", "sample_size_n", "confidence") for r in (1, 2, 3)] + [final[rid]["agreement"]] for rid in rids]
    write_sheet(ws, ccols, rows, {c: 22 for c in ccols})

    # Human_verification: tiered sample (see METHODS_AND_VERIFICATION.md)
    rng = random.Random(SEED)
    tier = {}
    for rid in rids:
        f = final[rid]
        if f["flag"] == "Yes":
            tier[rid] = "1 — flagged"
        elif f["study_type"] in COMPARATIVE:
            tier[rid] = "2 — comparative study (drives evidence-type counts)"
    rest = [rid for rid in rids if rid not in tier]
    by_type = {}
    for rid in rest:
        by_type.setdefault(final[rid]["study_type"], []).append(rid)
    for t, ids in by_type.items():
        k = max(1, round(0.2 * len(ids)))
        for rid in rng.sample(ids, k):
            tier[rid] = "3 — 20% random audit of unflagged"
    ws = wb.create_sheet("Human_verification")
    hcols = ["tier", "ref_id", "recs", "citation", "paper_file", "pmid", "evidence_basis", "study_type",
             "severity_adjustment", "sibling_analysis_any", "sample_size_n", "sample_size_unit", "flag_reasons",
             "pmid_correct (Y/N)", "study_type_correct (Y/N)", "corrected_study_type", "severity_correct (Y/N)",
             "corrected_severity", "sample_size_correct (Y/N)", "corrected_sample_size", "reviewer_initials",
             "minutes_spent", "human_notes"]
    rows = []
    for rid in sorted(tier, key=lambda x: (tier[x], sort_key(x))):
        f, c = final[rid], cits[rid]
        rows.append([tier[rid], rid, c["recs"], c["citation"], c["paper_file"], c.get("pmid"), f["evidence_basis"],
                     f["study_type"], f["severity_adjustment"], f["sibling_analysis_any"], f["sample_size_n"],
                     f["sample_size_unit"], f["flag_reasons"]] + [None] * 10)
    write_sheet(ws, hcols, rows, WIDE, input_cols=range(13, 23))

    # No_full_text
    ws = wb.create_sheet("No_full_text")
    ncols = ["ref_id", "recs", "citation", "pmid", "pmcid", "full_text_status"]
    rows = [[cits[r].get(k) for k in ncols] for r in rids if not cits[r]["full_text_available"]]
    write_sheet(ws, ncols, rows, WIDE)

    wb.save(OUT / "Compiled_evidence_coding.xlsx")
    (WORK / "final_codes.json").write_text(json.dumps(final, indent=1, ensure_ascii=False))
    print("tiers:", Counter(tier.values()))
    print("flagged:", sum(f["flag"] == "Yes" for f in final.values()), "of", len(final))
    print("final study types:", Counter(f["study_type"] for f in final.values()))


if __name__ == "__main__":
    main()
